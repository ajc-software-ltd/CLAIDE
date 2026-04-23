#!/usr/bin/env python3
from __future__ import annotations

import argparse
import base64
import importlib
import io
import json
import os
import tempfile
import traceback
import zipfile
from typing import Dict, List, Optional, Set, Tuple

PROTOCOL_VERSION = "1"

ALLOWED_PROVIDERS = {
    "system": {"probe"},
    "smoke": {"ping"},
    "mammoth": {"docx_to_html"},
    "mammoth_semantic": {"export_html", "extract_raw_text"},
    "docxcompose": {"compose_append"},
    "lxml": {"xpath_query", "xslt_transform"},
    "schematron": {"validate_part"},
    "ocr": {"extract_text"},
    "pypdf": {"extract_text"},
    "pdfminer": {"extract_text", "extract_layout"},
    "docxtpl": {"render_template"},
    "python_docx": {"style_audit"},
}

PROVIDER_CAPABILITIES = {
    "smoke": ["ping"],
    "mammoth": ["docx_to_html"],
    "mammoth_semantic": ["export_html", "extract_raw_text"],
    "docxcompose": ["compose_append"],
    "lxml": ["xpath_query", "xslt_transform"],
    "schematron": ["validate_part"],
    "ocr": ["extract_text"],
    "pypdf": ["extract_text"],
    "pdfminer": ["extract_text", "extract_layout"],
    "docxtpl": ["render_template"],
    "python_docx": ["style_audit"],
}

LXML_ALLOWED_PARTS = {
    "word/document.xml",
    "word/styles.xml",
    "word/numbering.xml",
    "_rels/.rels",
    "word/_rels/document.xml.rels",
    "[Content_Types].xml",
}

SCHEMATRON_ALLOWED_PARTS = {
    "word/document.xml",
    "word/styles.xml",
    "word/numbering.xml",
    "_rels/.rels",
    "word/_rels/document.xml.rels",
    "[Content_Types].xml",
}


def escape(value: str) -> str:
    return value.replace("\\", "\\\\").replace("\n", "\\n")


def unescape(value: str) -> str:
    return value.replace("\\n", "\n").replace("\\\\", "\\")


def read_kv(path: str) -> Dict[str, str]:
    out: Dict[str, str] = {}
    with open(path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.rstrip("\n")
            if "=" not in line:
                continue
            k, v = line.split("=", 1)
            out[k.strip()] = v
    return out


def write_kv(path: str, values: Dict[str, str]) -> None:
    values = {"protocol_version": PROTOCOL_VERSION, **values}
    with open(path, "w", encoding="utf-8") as f:
        for k, v in values.items():
            f.write(f"{k}={escape(str(v))}\n")


def decode_payload(values: Dict[str, str]) -> str:
    encoded = values.get("payload_b64", "")
    if not encoded:
        return ""
    try:
        return base64.b64decode(encoded.encode("utf-8")).decode("utf-8", errors="replace")
    except Exception:
        return ""


def encode_text(text: str) -> str:
    return base64.b64encode(text.encode("utf-8")).decode("utf-8")


def error(code: str, message: str, debug: str = "") -> Dict[str, str]:
    out = {"code": code, "message": message, "provenance": "python_provider"}
    if debug:
        out["debug_detail"] = debug
    return out


def module_version(module_name: str) -> Tuple[bool, str, str]:
    try:
        mod = importlib.import_module(module_name)
        version = getattr(mod, "__version__", "unknown")
        return True, str(version), ""
    except Exception as ex:
        return False, "", str(ex)


def pack_providers(providers: List[Tuple[str, bool, str, List[str], str]]) -> str:
    chunks = []
    for name, avail, version, capabilities, message in providers:
        caps = ",".join(capabilities)
        chunks.append(f"{escape(name)}|{'1' if avail else '0'}|{escape(version)}|{escape(caps)}|{escape(message)}")
    return ";".join(chunks)


def mammoth_messages_to_warning_text(messages: List[object]) -> str:
    lines = []
    for m in messages:
        text = getattr(m, "message", str(m))
        level = getattr(m, "type", "warning")
        lines.append(f"{level}:{text}")
    return "\n".join(lines)


def provider_probe() -> Dict[str, str]:
    providers: List[Tuple[str, bool, str, List[str], str]] = []

    providers.append(("smoke", True, "builtin", PROVIDER_CAPABILITIES["smoke"], "available"))

    ok, ver, msg = module_version("mammoth")
    providers.append(("mammoth", ok, ver, PROVIDER_CAPABILITIES["mammoth"], "available" if ok else msg))
    providers.append(("mammoth_semantic", ok, ver, PROVIDER_CAPABILITIES["mammoth_semantic"], "available" if ok else msg))

    okc, verc, msgc = module_version("docxcompose")
    providers.append(("docxcompose", okc, verc, PROVIDER_CAPABILITIES["docxcompose"], "available" if okc else msgc))

    okl, verl, msgl = module_version("lxml")
    providers.append(("lxml", okl, verl, PROVIDER_CAPABILITIES["lxml"], "available" if okl else msgl))

    if okl:
        try:
            from lxml import isoschematron  # type: ignore

            _ = isoschematron.Schematron
            providers.append(("schematron", True, verl, PROVIDER_CAPABILITIES["schematron"], "available"))
        except Exception as ex:
            providers.append(("schematron", False, verl, PROVIDER_CAPABILITIES["schematron"], str(ex)))
    else:
        providers.append(("schematron", False, "", PROVIDER_CAPABILITIES["schematron"], msgl))

    oko, vero, msgo = module_version("pytesseract")
    if oko:
        try:
            import pytesseract  # type: ignore

            _ = pytesseract.get_tesseract_version()
            providers.append(("ocr", True, vero, PROVIDER_CAPABILITIES["ocr"], "available"))
        except Exception as ex:
            providers.append(("ocr", False, vero, PROVIDER_CAPABILITIES["ocr"], str(ex)))
    else:
        providers.append(("ocr", False, "", PROVIDER_CAPABILITIES["ocr"], msgo))

    okp, verp, msgp = module_version("pypdf")
    providers.append(("pypdf", okp, verp, PROVIDER_CAPABILITIES["pypdf"], "available" if okp else msgp))

    okm, verm, msgm = module_version("pdfminer")
    providers.append(("pdfminer", okm, verm, PROVIDER_CAPABILITIES["pdfminer"], "available" if okm else msgm))

    okt, vert, msqt = module_version("docxtpl")
    providers.append(("docxtpl", okt, vert, PROVIDER_CAPABILITIES["docxtpl"], "available" if okt else msqt))

    okd, verd, msgd = module_version("docx")
    providers.append(("python_docx", okd, verd, PROVIDER_CAPABILITIES["python_docx"], "available" if okd else msgd))

    return {
        "code": "ok",
        "message": "provider probe complete",
        "provenance": "python_provider",
        "providers": pack_providers(providers),
    }


def provider_smoke_ping() -> Dict[str, str]:
    return {
        "code": "ok",
        "message": "smoke provider alive",
        "provenance": "python_provider",
        "text_b64": encode_text("pong"),
    }


def provider_mammoth_docx_to_html(input_path: str) -> Dict[str, str]:
    try:
        import mammoth  # type: ignore
    except Exception as ex:
        return error("provider_unavailable", f"mammoth unavailable: {ex}")

    if not input_path or not os.path.exists(input_path):
        return error("invalid_request", "input_path for mammoth is required")

    with open(input_path, "rb") as docx_file:
        result = mammoth.convert_to_html(docx_file)
    return {
        "code": "ok",
        "message": "mammoth conversion complete",
        "provider_version": getattr(mammoth, "__version__", "unknown"),
        "provenance": "python_provider",
        "text_b64": encode_text(result.value),
    }


def provider_mammoth_semantic_export_html(input_path: str, payload: str) -> Dict[str, str]:
    try:
        import mammoth  # type: ignore
    except Exception as ex:
        return error("provider_unavailable", f"mammoth unavailable: {ex}")

    if not input_path or not os.path.exists(input_path):
        return error("invalid_request", "input_path for mammoth_semantic.export_html is required")

    style_map = None
    include_default_style_map = True
    if payload:
        try:
            params = json.loads(payload)
            if not isinstance(params, dict):
                return error("invalid_request", "payload must be a JSON object")
            style_map = params.get("style_map")
            include_default_style_map = bool(params.get("include_default_style_map", True))
        except Exception as ex:
            return error("invalid_request", f"invalid payload JSON: {ex}")

    kwargs = {"include_default_style_map": include_default_style_map}
    if style_map:
        kwargs["style_map"] = style_map

    with open(input_path, "rb") as docx_file:
        result = mammoth.convert_to_html(docx_file, **kwargs)

    warnings = mammoth_messages_to_warning_text(result.messages)

    out = {
        "code": "ok",
        "message": "mammoth semantic html export complete",
        "provider_version": getattr(mammoth, "__version__", "unknown"),
        "provenance": "python_provider",
        "text_b64": encode_text(result.value),
    }
    if warnings:
        out["warnings_b64"] = encode_text(warnings)
    return out


def provider_mammoth_semantic_extract_text(input_path: str) -> Dict[str, str]:
    try:
        import mammoth  # type: ignore
    except Exception as ex:
        return error("provider_unavailable", f"mammoth unavailable: {ex}")

    if not input_path or not os.path.exists(input_path):
        return error("invalid_request", "input_path for mammoth_semantic.extract_raw_text is required")

    with open(input_path, "rb") as docx_file:
        result = mammoth.extract_raw_text(docx_file)

    warnings = mammoth_messages_to_warning_text(getattr(result, "messages", []))

    out = {
        "code": "ok",
        "message": "mammoth semantic raw text extraction complete",
        "provider_version": getattr(mammoth, "__version__", "unknown"),
        "provenance": "python_provider",
        "text_b64": encode_text(result.value),
    }
    if warnings:
        out["warnings_b64"] = encode_text(warnings)
    return out


def provider_docxcompose_append(base_path: str, append_path: str, output_path: str) -> Dict[str, str]:
    try:
        from docx import Document  # type: ignore
        from docxcompose.composer import Composer  # type: ignore
    except Exception as ex:
        return error("provider_unavailable", f"docxcompose unavailable: {ex}")

    if not base_path or not append_path or not output_path:
        return error("invalid_request", "base input_path, append payload, and output_path are required")

    if not os.path.exists(base_path):
        return error("invalid_request", "base DOCX path not found")
    if not os.path.exists(append_path):
        return error("invalid_request", "append DOCX path not found")

    base = Document(base_path)
    composer = Composer(base)
    composer.append(Document(append_path))
    composer.save(output_path)

    return {
        "code": "ok",
        "message": "docxcompose append complete",
        "provenance": "python_provider",
        "output_path": output_path,
    }


def provider_lxml_xpath_query(input_path: str, xpath_expr: str) -> Dict[str, str]:
    return provider_lxml_xpath_query_structured(input_path, xpath_expr)


def _read_docx_part_xml_allowlisted(docx_path: str, part_path: str, allowed_parts: Set[str]) -> Tuple[bool, str, str]:
    if not docx_path or not os.path.exists(docx_path):
        return False, "", "input_path is required"
    if part_path not in allowed_parts:
        return False, "", f"part not allowed: {part_path}"

    try:
        with zipfile.ZipFile(docx_path, "r") as zf:
            with zf.open(part_path, "r") as part:
                raw = part.read()
    except KeyError:
        return False, "", f"requested part not found in DOCX: {part_path}"
    except Exception as ex:
        return False, "", f"failed to read DOCX part: {ex}"

    try:
        return True, raw.decode("utf-8"), ""
    except Exception as ex:
        return False, "", f"failed to decode XML part as utf-8: {ex}"


def _lxml_version() -> str:
    try:
        import lxml  # type: ignore

        return str(getattr(lxml, "__version__", "unknown"))
    except Exception:
        return "unknown"


def provider_lxml_xpath_query_structured(input_path: str, payload_json: str) -> Dict[str, str]:
    try:
        from lxml import etree  # type: ignore
    except Exception as ex:
        return error("provider_unavailable", f"lxml unavailable: {ex}")

    if not payload_json:
        return error("invalid_request", "lxml xpath_query payload is required")

    try:
        payload = json.loads(payload_json)
    except Exception as ex:
        return error("invalid_request", f"invalid lxml xpath payload JSON: {ex}")
    if not isinstance(payload, dict):
        return error("invalid_request", "lxml xpath payload must be a JSON object")

    part = str(payload.get("part", ""))
    xpath_expr = str(payload.get("xpath", ""))
    mode = str(payload.get("mode", "xpath"))
    namespaces = payload.get("namespaces", {})

    if not part or not xpath_expr:
        return error("invalid_request", "lxml xpath requires payload.part and payload.xpath")
    if mode not in {"xpath", "compiled_xpath", "evaluator"}:
        return error("invalid_request", "unsupported xpath mode; expected xpath|compiled_xpath|evaluator")
    if not isinstance(namespaces, dict):
        return error("invalid_request", "payload.namespaces must be an object when provided")

    ok, xml_text, read_error = _read_docx_part_xml_allowlisted(input_path, part, LXML_ALLOWED_PARTS)
    if not ok:
        return error("invalid_request", read_error)

    try:
        root = etree.fromstring(xml_text.encode("utf-8"))
    except Exception as ex:
        return error("execution_failed", f"failed to parse selected XML part: {ex}")

    try:
        if mode == "compiled_xpath":
            evaluator = etree.XPath(xpath_expr, namespaces=namespaces)
            result = evaluator(root)
        elif mode == "evaluator":
            evaluator = etree.XPathElementEvaluator(root, namespaces=namespaces)
            result = evaluator(xpath_expr)
        else:
            result = root.xpath(xpath_expr, namespaces=namespaces)
    except Exception as ex:
        return error("invalid_request", f"invalid xpath expression or evaluation failed: {ex}")

    normalized_items = []
    for item in result if isinstance(result, list) else [result]:
        item_type = type(item).__name__
        item_value = ""
        item_path = ""
        if hasattr(item, "getroottree") and hasattr(item, "tag"):
            try:
                item_path = item.getroottree().getpath(item)
            except Exception:
                item_path = ""
            try:
                item_value = etree.tostring(item, encoding="unicode")
            except Exception:
                item_value = str(item)
        else:
            item_value = str(item)
        normalized_items.append({"type": item_type, "value": item_value, "path": item_path})

    result_payload = {
        "provider": "lxml",
        "provider_version": _lxml_version(),
        "operation": "xpath_query",
        "selected_part": part,
        "xpath_mode": mode,
        "success": True,
        "warnings": [],
        "errors": [],
        "provenance": "python_provider",
        "result": normalized_items,
    }

    return {
        "code": "ok",
        "message": "xpath query complete",
        "provider_version": _lxml_version(),
        "provenance": "python_provider",
        "text_b64": encode_text(json.dumps(result_payload, ensure_ascii=False)),
    }


def provider_lxml_xslt_transform(input_path: str, xslt_text: str) -> Dict[str, str]:
    return provider_lxml_xslt_transform_structured(input_path, xslt_text)


def provider_lxml_xslt_transform_structured(input_path: str, payload_json: str) -> Dict[str, str]:
    try:
        from lxml import etree  # type: ignore
    except Exception as ex:
        return error("provider_unavailable", f"lxml unavailable: {ex}")

    if not payload_json:
        return error("invalid_request", "lxml xslt_transform payload is required")

    try:
        payload = json.loads(payload_json)
    except Exception as ex:
        return error("invalid_request", f"invalid lxml xslt payload JSON: {ex}")
    if not isinstance(payload, dict):
        return error("invalid_request", "lxml xslt payload must be a JSON object")

    part = str(payload.get("part", ""))
    xslt_text = str(payload.get("xslt", ""))
    params = payload.get("params", {})
    output_mode = str(payload.get("output_mode", "xml"))
    if not part or not xslt_text:
        return error("invalid_request", "lxml xslt requires payload.part and payload.xslt")
    if output_mode not in {"xml", "text"}:
        return error("invalid_request", "unsupported output_mode; expected xml|text")
    if not isinstance(params, dict):
        return error("invalid_request", "payload.params must be an object when provided")

    ok, xml_text, read_error = _read_docx_part_xml_allowlisted(input_path, part, LXML_ALLOWED_PARTS)
    if not ok:
        return error("invalid_request", read_error)

    try:
        xml_doc = etree.parse(io.BytesIO(xml_text.encode("utf-8")))
    except Exception as ex:
        return error("execution_failed", f"failed to parse selected XML part: {ex}")

    try:
        xslt_doc = etree.parse(io.BytesIO(xslt_text.encode("utf-8")))
        transform = etree.XSLT(xslt_doc)
    except Exception as ex:
        return error("invalid_request", f"invalid xslt stylesheet: {ex}")

    try:
        xslt_params = {k: etree.XSLT.strparam(str(v)) for k, v in params.items()}
        output = transform(xml_doc, **xslt_params)
    except Exception as ex:
        return error("execution_failed", f"xslt transform failed: {ex}")

    transformed = str(output) if output_mode == "xml" else str(output)
    result_payload = {
        "provider": "lxml",
        "provider_version": _lxml_version(),
        "operation": "xslt_transform",
        "selected_part": part,
        "output_mode": output_mode,
        "success": True,
        "warnings": [],
        "errors": [],
        "provenance": "python_provider",
        "result": transformed,
    }

    return {
        "code": "ok",
        "message": "xslt transform complete",
        "provider_version": _lxml_version(),
        "provenance": "python_provider",
        "text_b64": encode_text(json.dumps(result_payload, ensure_ascii=False)),
    }


def provider_schematron_validate_part(input_path: str, payload_json: str) -> Dict[str, str]:
    try:
        from lxml import etree  # type: ignore
        from lxml import isoschematron  # type: ignore
    except Exception as ex:
        return error("provider_unavailable", f"lxml/isoschematron unavailable: {ex}")

    if not payload_json:
        return error("invalid_request", "schematron validate_part payload is required")
    try:
        payload = json.loads(payload_json)
    except Exception as ex:
        return error("invalid_request", f"invalid schematron payload JSON: {ex}")
    if not isinstance(payload, dict):
        return error("invalid_request", "schematron payload must be a JSON object")

    part = str(payload.get("part", ""))
    schema_text = payload.get("schema_text")
    schema_path = payload.get("schema_path")
    phase = payload.get("phase")
    store_report = bool(payload.get("store_report", True))

    if not part:
        return error("invalid_request", "schematron validate_part requires payload.part")
    if bool(schema_text) == bool(schema_path):
        return error("invalid_request", "provide exactly one of payload.schema_text or payload.schema_path")
    if schema_text is not None and not isinstance(schema_text, str):
        return error("invalid_request", "payload.schema_text must be a string when provided")
    if schema_path is not None and not isinstance(schema_path, str):
        return error("invalid_request", "payload.schema_path must be a string when provided")
    if phase is not None and not isinstance(phase, str):
        return error("invalid_request", "payload.phase must be a string when provided")

    ok, xml_text, read_error = _read_docx_part_xml_allowlisted(input_path, part, SCHEMATRON_ALLOWED_PARTS)
    if not ok:
        return error("invalid_request", read_error)

    try:
        target_doc = etree.parse(io.BytesIO(xml_text.encode("utf-8")))
    except Exception as ex:
        return error("execution_failed", f"failed to parse selected XML part: {ex}")

    try:
        if schema_text:
            schema_doc = etree.parse(io.BytesIO(str(schema_text).encode("utf-8")))
            schema_source = "text"
        else:
            if not os.path.exists(str(schema_path)):
                return error("invalid_request", f"schematron schema file not found: {schema_path}")
            schema_doc = etree.parse(str(schema_path))
            schema_source = "file"
    except Exception as ex:
        return error("invalid_request", f"invalid schematron schema input: {ex}")

    try:
        validator = isoschematron.Schematron(schema_doc, store_report=store_report, phase=phase)
    except Exception as ex:
        return error("invalid_request", f"failed to compile schematron schema: {ex}")

    try:
        valid = bool(validator.validate(target_doc))
    except Exception as ex:
        return error("execution_failed", f"schematron validation failed: {ex}")

    failed_asserts: List[Dict[str, str]] = []
    report_entries: List[Dict[str, str]] = []
    report_xml = ""
    try:
        report = getattr(validator, "validation_report", None)
        if report is not None:
            report_xml = etree.tostring(report, encoding="unicode")
            ns = {"svrl": "http://purl.oclc.org/dsdl/svrl"}
            for node in report.xpath("//svrl:failed-assert", namespaces=ns):
                text_nodes = node.xpath("./svrl:text/text()", namespaces=ns)
                failed_asserts.append(
                    {
                        "location": str(node.get("location", "")),
                        "test": str(node.get("test", "")),
                        "text": str(text_nodes[0]) if text_nodes else "",
                    }
                )
            for node in report.xpath("//svrl:successful-report", namespaces=ns):
                text_nodes = node.xpath("./svrl:text/text()", namespaces=ns)
                report_entries.append(
                    {
                        "location": str(node.get("location", "")),
                        "test": str(node.get("test", "")),
                        "text": str(text_nodes[0]) if text_nodes else "",
                    }
                )
    except Exception:
        pass

    result_payload = {
        "provider": "schematron",
        "provider_version": _lxml_version(),
        "operation": "validate_part",
        "selected_part": part,
        "schema_source": schema_source,
        "phase": phase if phase is not None else "",
        "success": True,
        "valid": valid,
        "failed_asserts": failed_asserts,
        "reports": report_entries,
        "report_xml": report_xml if store_report else "",
        "warnings": [],
        "errors": [],
        "provenance": "python_provider",
    }

    return {
        "code": "ok",
        "message": "schematron validation complete",
        "provider_version": _lxml_version(),
        "provenance": "python_provider",
        "text_b64": encode_text(json.dumps(result_payload, ensure_ascii=False)),
    }


def provider_ocr_extract_text(input_path: str, payload_json: str) -> Dict[str, str]:
    try:
        import pytesseract  # type: ignore
    except Exception as ex:
        return error("provider_unavailable", f"pytesseract unavailable: {ex}")

    try:
        from PIL import Image  # type: ignore
    except Exception as ex:
        return error("provider_unavailable", f"Pillow unavailable: {ex}")

    try:
        _ = pytesseract.get_tesseract_version()
    except Exception as ex:
        return error("provider_unavailable", f"tesseract runtime unavailable: {ex}")

    payload: Dict[str, object] = {}
    if payload_json:
        try:
            parsed = json.loads(payload_json)
        except Exception as ex:
            return error("invalid_request", f"invalid ocr payload JSON: {ex}")
        if not isinstance(parsed, dict):
            return error("invalid_request", "ocr payload must be a JSON object")
        payload = parsed

    lang = str(payload.get("lang", "eng"))
    psm = payload.get("psm")
    image_b64 = payload.get("image_b64")
    image_format = str(payload.get("image_format", "png")).lower()

    if psm is not None and psm not in {3, 6, 11}:
        return error("invalid_request", "unsupported psm value; allowed values are 3, 6, 11")
    if image_b64 is not None and not isinstance(image_b64, str):
        return error("invalid_request", "image_b64 must be a base64 string when provided")

    temp_image_path = ""
    target_image_path = input_path
    if image_b64:
        try:
            image_bytes = base64.b64decode(image_b64.encode("utf-8"))
        except Exception as ex:
            return error("invalid_request", f"invalid image_b64 payload: {ex}")
        suffix = f".{image_format}" if image_format in {"png", "jpg", "jpeg", "bmp", "tif", "tiff"} else ".png"
        with tempfile.NamedTemporaryFile(prefix="minidocx_ocr_", suffix=suffix, delete=False) as tmp:
            tmp.write(image_bytes)
            temp_image_path = tmp.name
            target_image_path = temp_image_path

    if not target_image_path or not os.path.exists(target_image_path):
        if temp_image_path:
            try:
                os.remove(temp_image_path)
            except Exception:
                pass
        return error("invalid_request", "ocr requires input_path or payload.image_b64 with valid image data")

    config_chunks = []
    if psm is not None:
        config_chunks.extend(["--psm", str(psm)])
    config = " ".join(config_chunks).strip()

    try:
        with Image.open(target_image_path) as image:
            text = pytesseract.image_to_string(image, lang=lang, config=config)
    except pytesseract.TesseractNotFoundError as ex:  # type: ignore[attr-defined]
        return error("provider_unavailable", f"tesseract executable not found: {ex}")
    except Exception as ex:
        message = str(ex)
        if "Error opening data file" in message or "Failed loading language" in message:
            return error("invalid_request", f"ocr language data unavailable: {message}")
        return error("execution_failed", f"ocr extraction failed: {message}")
    finally:
        if temp_image_path:
            try:
                os.remove(temp_image_path)
            except Exception:
                pass

    result_payload = {
        "provider": "ocr",
        "provider_version": getattr(pytesseract, "__version__", "unknown"),
        "engine": "tesseract",
        "operation": "extract_text",
        "success": True,
        "warnings": [],
        "errors": [],
        "provenance": "python_provider",
        "options": {"lang": lang, "psm": psm},
        "text": text,
    }

    return {
        "code": "ok",
        "message": "ocr extraction complete",
        "provider_version": getattr(pytesseract, "__version__", "unknown"),
        "provenance": "python_provider",
        "text_b64": encode_text(json.dumps(result_payload, ensure_ascii=False)),
    }


def provider_pypdf_extract_text(input_path: str, payload_json: str) -> Dict[str, str]:
    try:
        from pypdf import PdfReader  # type: ignore
        import pypdf  # type: ignore
    except Exception as ex:
        return error("provider_unavailable", f"pypdf unavailable: {ex}")

    payload: Dict[str, object] = {}
    if payload_json:
        try:
            parsed = json.loads(payload_json)
        except Exception as ex:
            return error("invalid_request", f"invalid pypdf payload JSON: {ex}")
        if not isinstance(parsed, dict):
            return error("invalid_request", "pypdf payload must be a JSON object")
        payload = parsed

    extraction_mode = str(payload.get("mode", "plain")).strip().lower()
    pdf_b64 = payload.get("pdf_b64")
    if extraction_mode not in {"plain", "layout"}:
        return error("invalid_request", "pypdf mode must be one of: plain, layout")
    if pdf_b64 is not None and not isinstance(pdf_b64, str):
        return error("invalid_request", "payload.pdf_b64 must be a base64 string when provided")

    pdf_path = input_path
    temp_pdf_path = ""
    if pdf_b64:
        try:
            pdf_bytes = base64.b64decode(pdf_b64.encode("utf-8"))
        except Exception as ex:
            return error("invalid_request", f"invalid payload.pdf_b64 value: {ex}")
        with tempfile.NamedTemporaryFile(prefix="minidocx_pdf_", suffix=".pdf", delete=False) as tmp:
            tmp.write(pdf_bytes)
            temp_pdf_path = tmp.name
            pdf_path = temp_pdf_path

    if not pdf_path or not os.path.exists(pdf_path):
        if temp_pdf_path:
            try:
                os.remove(temp_pdf_path)
            except Exception:
                pass
        return error("invalid_request", "pypdf extract_text requires input_path or payload.pdf_b64")

    warnings: List[str] = []
    extracted_pages: List[str] = []
    try:
        reader = PdfReader(pdf_path)
        for page in reader.pages:
            if extraction_mode == "layout":
                text = page.extract_text(extraction_mode="layout")
            else:
                text = page.extract_text()
            extracted_pages.append(text or "")
    except Exception as ex:
        return error("execution_failed", f"pypdf extraction failed: {ex}")
    finally:
        if temp_pdf_path:
            try:
                os.remove(temp_pdf_path)
            except Exception:
                pass

    output_text = "\n".join(extracted_pages)
    if not output_text.strip():
        warnings.append("PDF text extraction returned empty/minimal text; scanned/image-only PDFs may require OCR.")

    result_payload = {
        "provider": "pypdf",
        "provider_version": getattr(pypdf, "__version__", "unknown"),
        "operation": "extract_text",
        "extraction_mode": extraction_mode,
        "page_count": len(extracted_pages),
        "success": True,
        "warnings": warnings,
        "errors": [],
        "provenance": "python_provider",
        "text": output_text,
    }
    out = {
        "code": "ok",
        "message": "pypdf extraction complete",
        "provider_version": getattr(pypdf, "__version__", "unknown"),
        "provenance": "python_provider",
        "text_b64": encode_text(json.dumps(result_payload, ensure_ascii=False)),
    }
    if warnings:
        out["warnings_b64"] = encode_text("\n".join(warnings))
    return out


def _decode_pdf_source(input_path: str, payload: Dict[str, object]) -> Tuple[str, str, Optional[Dict[str, str]]]:
    pdf_b64 = payload.get("pdf_b64")
    if pdf_b64 is not None and not isinstance(pdf_b64, str):
        return "", "", error("invalid_request", "payload.pdf_b64 must be a base64 string when provided")

    pdf_path = input_path
    temp_pdf_path = ""
    if pdf_b64:
        try:
            pdf_bytes = base64.b64decode(pdf_b64.encode("utf-8"))
        except Exception as ex:
            return "", "", error("invalid_request", f"invalid payload.pdf_b64 value: {ex}")
        with tempfile.NamedTemporaryFile(prefix="minidocx_pdfminer_", suffix=".pdf", delete=False) as tmp:
            tmp.write(pdf_bytes)
            temp_pdf_path = tmp.name
            pdf_path = temp_pdf_path

    if not pdf_path or not os.path.exists(pdf_path):
        if temp_pdf_path:
            try:
                os.remove(temp_pdf_path)
            except Exception:
                pass
        return "", "", error("invalid_request", "pdfminer operation requires input_path or payload.pdf_b64")
    return pdf_path, temp_pdf_path, None


def _cleanup_temp_pdf(path: str) -> None:
    if path:
        try:
            os.remove(path)
        except Exception:
            pass


def _parse_laparams_subset(
    payload: Dict[str, object],
    laparams_cls: object,
) -> Tuple[Optional[object], Dict[str, object], Optional[Dict[str, str]]]:
    raw = payload.get("laparams")
    if raw is None:
        return None, {}, None
    if not isinstance(raw, dict):
        return None, {}, error("invalid_request", "payload.laparams must be an object when provided")

    allowed_keys = {"char_margin", "word_margin", "line_margin", "boxes_flow"}
    for key in raw.keys():
        if key not in allowed_keys:
            return None, {}, error("invalid_request", f"unsupported laparams key: {key}")

    kwargs: Dict[str, object] = {}
    for key in allowed_keys:
        if key in raw:
            value = raw[key]
            if not isinstance(value, (int, float)):
                return None, {}, error("invalid_request", f"laparams.{key} must be numeric")
            kwargs[key] = float(value)
    return laparams_cls(**kwargs), kwargs, None


def provider_pdfminer_extract_text(input_path: str, payload_json: str) -> Dict[str, str]:
    try:
        from pdfminer.high_level import extract_text  # type: ignore
        from pdfminer.layout import LAParams  # type: ignore
        import pdfminer  # type: ignore
    except Exception as ex:
        return error("provider_unavailable", f"pdfminer unavailable: {ex}")

    payload: Dict[str, object] = {}
    if payload_json:
        try:
            parsed = json.loads(payload_json)
        except Exception as ex:
            return error("invalid_request", f"invalid pdfminer payload JSON: {ex}")
        if not isinstance(parsed, dict):
            return error("invalid_request", "pdfminer payload must be a JSON object")
        payload = parsed

    page_numbers_raw = payload.get("page_numbers")
    page_numbers = None
    if page_numbers_raw is not None:
        if not isinstance(page_numbers_raw, list) or not all(isinstance(v, int) and v >= 0 for v in page_numbers_raw):
            return error("invalid_request", "payload.page_numbers must be an array of non-negative integers")
        page_numbers = [int(v) for v in page_numbers_raw]

    laparams, laparams_kwargs, lap_err = _parse_laparams_subset(payload, LAParams)
    if lap_err is not None:
        return lap_err

    pdf_path, temp_pdf_path, source_err = _decode_pdf_source(input_path, payload)
    if source_err is not None:
        return source_err

    warnings: List[str] = []
    try:
        text = extract_text(pdf_path, page_numbers=page_numbers, laparams=laparams)
    except Exception as ex:
        _cleanup_temp_pdf(temp_pdf_path)
        return error("execution_failed", f"pdfminer extract_text failed: {ex}")
    _cleanup_temp_pdf(temp_pdf_path)

    if not text.strip():
        warnings.append("PDF text extraction returned empty/minimal text; scanned/image-only PDFs may require OCR.")

    result_payload = {
        "provider": "pdfminer",
        "provider_version": getattr(pdfminer, "__version__", "unknown"),
        "operation": "extract_text",
        "page_numbers": page_numbers if page_numbers is not None else [],
        "laparams": laparams_kwargs,
        "success": True,
        "warnings": warnings,
        "errors": [],
        "provenance": "python_provider",
        "text": text,
    }

    out = {
        "code": "ok",
        "message": "pdfminer text extraction complete",
        "provider_version": getattr(pdfminer, "__version__", "unknown"),
        "provenance": "python_provider",
        "text_b64": encode_text(json.dumps(result_payload, ensure_ascii=False)),
    }
    if warnings:
        out["warnings_b64"] = encode_text("\n".join(warnings))
    return out


def provider_pdfminer_extract_layout(input_path: str, payload_json: str) -> Dict[str, str]:
    try:
        from pdfminer.high_level import extract_pages  # type: ignore
        from pdfminer.layout import LAParams, LTChar, LTTextBox, LTTextLine  # type: ignore
        import pdfminer  # type: ignore
    except Exception as ex:
        return error("provider_unavailable", f"pdfminer unavailable: {ex}")

    payload: Dict[str, object] = {}
    if payload_json:
        try:
            parsed = json.loads(payload_json)
        except Exception as ex:
            return error("invalid_request", f"invalid pdfminer payload JSON: {ex}")
        if not isinstance(parsed, dict):
            return error("invalid_request", "pdfminer payload must be a JSON object")
        payload = parsed

    page_numbers_raw = payload.get("page_numbers")
    page_numbers = None
    if page_numbers_raw is not None:
        if not isinstance(page_numbers_raw, list) or not all(isinstance(v, int) and v >= 0 for v in page_numbers_raw):
            return error("invalid_request", "payload.page_numbers must be an array of non-negative integers")
        page_numbers = [int(v) for v in page_numbers_raw]

    laparams, laparams_kwargs, lap_err = _parse_laparams_subset(payload, LAParams)
    if lap_err is not None:
        return lap_err

    pdf_path, temp_pdf_path, source_err = _decode_pdf_source(input_path, payload)
    if source_err is not None:
        return source_err

    page_summaries: List[Dict[str, object]] = []
    warnings: List[str] = []
    try:
        for page_index, page_layout in enumerate(extract_pages(pdf_path, page_numbers=page_numbers, laparams=laparams)):
            text_boxes = []
            line_count = 0
            char_count = 0
            for element in page_layout:
                if isinstance(element, LTTextBox):
                    snippet = element.get_text().strip().replace("\n", " ")
                    if len(snippet) > 120:
                        snippet = snippet[:120]
                    text_boxes.append(
                        {
                            "bbox": [element.x0, element.y0, element.x1, element.y1],
                            "text": snippet,
                        }
                    )
                    for line in element:
                        if isinstance(line, LTTextLine):
                            line_count += 1
                            for ch in line:
                                if isinstance(ch, LTChar):
                                    char_count += 1
            page_summaries.append(
                {
                    "page_index": page_index,
                    "bbox": [page_layout.x0, page_layout.y0, page_layout.x1, page_layout.y1],
                    "text_box_count": len(text_boxes),
                    "text_line_count": line_count,
                    "char_count": char_count,
                    "text_boxes": text_boxes,
                }
            )
    except Exception as ex:
        _cleanup_temp_pdf(temp_pdf_path)
        return error("execution_failed", f"pdfminer extract_layout failed: {ex}")
    _cleanup_temp_pdf(temp_pdf_path)

    if not page_summaries:
        warnings.append("No pages were analyzed from the PDF input.")

    result_payload = {
        "provider": "pdfminer",
        "provider_version": getattr(pdfminer, "__version__", "unknown"),
        "operation": "extract_layout",
        "page_numbers": page_numbers if page_numbers is not None else [],
        "laparams": laparams_kwargs,
        "page_count": len(page_summaries),
        "success": True,
        "warnings": warnings,
        "errors": [],
        "provenance": "python_provider",
        "pages": page_summaries,
    }

    out = {
        "code": "ok",
        "message": "pdfminer layout extraction complete",
        "provider_version": getattr(pdfminer, "__version__", "unknown"),
        "provenance": "python_provider",
        "text_b64": encode_text(json.dumps(result_payload, ensure_ascii=False)),
    }
    if warnings:
        out["warnings_b64"] = encode_text("\n".join(warnings))
    return out


def provider_docxtpl_render_template(template_path: str, context_json: str, output_path: str) -> Dict[str, str]:
    try:
        from docxtpl import DocxTemplate  # type: ignore
    except Exception as ex:
        return error("provider_unavailable", f"docxtpl unavailable: {ex}")

    if not template_path or not output_path or not os.path.exists(template_path):
        return error("invalid_request", "template input_path and output_path are required")

    try:
        context = json.loads(context_json) if context_json else {}
        if not isinstance(context, dict):
            return error("invalid_request", "template context payload must be a JSON object")
    except Exception as ex:
        return error("invalid_request", f"invalid template context JSON: {ex}")

    tpl = DocxTemplate(template_path)
    tpl.render(context)
    tpl.save(output_path)

    return {
        "code": "ok",
        "message": "docxtpl render complete",
        "provenance": "python_provider",
        "output_path": output_path,
    }


def provider_python_docx_style_audit(input_path: str) -> Dict[str, str]:
    try:
        from docx import Document  # type: ignore
    except Exception as ex:
        return error("provider_unavailable", f"python-docx unavailable: {ex}")

    if not input_path or not os.path.exists(input_path):
        return error("invalid_request", "input_path is required for style audit")

    doc = Document(input_path)
    styles = []
    style_names = set()
    for style in doc.styles:
        style_names.add(style.name)
        styles.append({
            "name": style.name,
            "type": str(style.type),
            "builtin": bool(getattr(style, "builtin", False)),
        })

    missing_references = []
    for idx, para in enumerate(doc.paragraphs):
        if para.style is None:
            continue
        if para.style.name not in style_names:
            missing_references.append({"kind": "paragraph", "index": idx, "style": para.style.name})

    report = {
        "style_count": len(styles),
        "styles": styles,
        "missing_references": missing_references,
    }

    return {
        "code": "ok",
        "message": "python-docx style audit complete",
        "provenance": "python_provider",
        "text_b64": encode_text(json.dumps(report, ensure_ascii=False)),
    }


def validate_request(values: Dict[str, str]) -> Tuple[bool, Dict[str, str]]:
    required = ["protocol_version", "provider", "operation"]
    for field in required:
        if field not in values:
            return False, error("invalid_request", f"missing required field: {field}")

    if values.get("protocol_version") != PROTOCOL_VERSION:
        return False, error("protocol_mismatch", "request protocol version mismatch")

    provider = unescape(values.get("provider", ""))
    operation = unescape(values.get("operation", ""))

    if provider not in ALLOWED_PROVIDERS:
        return False, error("invalid_request", f"provider not allowed: {provider}")
    if operation not in ALLOWED_PROVIDERS[provider]:
        return False, error("invalid_request", f"operation not allowed: {provider}.{operation}")

    return True, {}


def run(values: Dict[str, str]) -> Dict[str, str]:
    ok, err = validate_request(values)
    if not ok:
        return err

    provider = unescape(values.get("provider", ""))
    operation = unescape(values.get("operation", ""))
    input_path = unescape(values.get("input_path", ""))
    output_path = unescape(values.get("output_path", ""))
    payload = decode_payload(values)

    if provider == "system" and operation == "probe":
        return provider_probe()
    if provider == "smoke" and operation == "ping":
        return provider_smoke_ping()
    if provider == "mammoth" and operation == "docx_to_html":
        return provider_mammoth_docx_to_html(input_path)
    if provider == "mammoth_semantic" and operation == "export_html":
        return provider_mammoth_semantic_export_html(input_path, payload)
    if provider == "mammoth_semantic" and operation == "extract_raw_text":
        return provider_mammoth_semantic_extract_text(input_path)
    if provider == "docxcompose" and operation == "compose_append":
        return provider_docxcompose_append(input_path, payload, output_path)
    if provider == "lxml" and operation == "xpath_query":
        return provider_lxml_xpath_query(input_path, payload)
    if provider == "lxml" and operation == "xslt_transform":
        return provider_lxml_xslt_transform(input_path, payload)
    if provider == "schematron" and operation == "validate_part":
        return provider_schematron_validate_part(input_path, payload)
    if provider == "ocr" and operation == "extract_text":
        return provider_ocr_extract_text(input_path, payload)
    if provider == "pypdf" and operation == "extract_text":
        return provider_pypdf_extract_text(input_path, payload)
    if provider == "pdfminer" and operation == "extract_text":
        return provider_pdfminer_extract_text(input_path, payload)
    if provider == "pdfminer" and operation == "extract_layout":
        return provider_pdfminer_extract_layout(input_path, payload)
    if provider == "docxtpl" and operation == "render_template":
        return provider_docxtpl_render_template(input_path, payload, output_path)
    if provider == "python_docx" and operation == "style_audit":
        return provider_python_docx_style_audit(input_path)

    return error("invalid_request", "unhandled provider operation")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--request", required=True)
    parser.add_argument("--response", required=True)
    args = parser.parse_args()

    try:
        request_values = {k: unescape(v) for k, v in read_kv(args.request).items()}
        response_values = run(request_values)
    except Exception as ex:
        response_values = error("execution_failed", f"worker exception: {ex}", traceback.format_exc())

    write_kv(args.response, response_values)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
